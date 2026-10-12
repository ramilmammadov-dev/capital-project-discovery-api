"""REST API (FastAPI TestClient). The pipeline is replaced by a canned result; storage and Excel are real."""

from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from capex.api.app import create_app
from capex.config import Settings
from capex.jobs import JobManager
from capex.llm import LLMUnavailableError
from capex.models import (
    ConsolidatedProject,
    DocStatus,
    DocType,
    DocumentCandidate,
    ExtractedRecord,
    FundingStatus,
    JobMetrics,
    JobResult,
    OrgProfile,
    Phase,
    SourceReference,
    UsageRecord,
)
from capex.pipeline import Pipeline
from capex.storage.cache import FileCache
from capex.storage.db import Database, JobRow

KEY = {"X-API-Key": "test-key"}
ORG = "City of Exampleville"
EXTRACT = "/api/v1/organizations/extract"


def make_documents() -> list[DocumentCandidate]:
    selected = DocumentCandidate(
        doc_id="doc_1", url="https://www.examplecity.gov/files/cip-2027.pdf", title="FY2027-2031 CIP",
        file_type="pdf", domain="examplecity.gov", is_official_domain=True, doc_type=DocType.CIP_PLAN,
        doc_status=DocStatus.ADOPTED, coverage_start_year=2027, coverage_end_year=2031, document_fiscal_year=2027,
        score=0.91, selected=True, selection_rank=1, selection_rationale="itemized project tables, adopted, newest",
    )
    other = DocumentCandidate(
        doc_id="doc_2", url="https://www.examplecity.gov/files/budget.pdf", title="Annual budget summary",
        file_type="pdf", domain="examplecity.gov", is_official_domain=True, score=0.31,
        rejection_reason="aggregate figures only",
    )
    return [selected, other]


def make_project(n: int, name: str, domain: str, budget: float, phase: Phase, funding: FundingStatus,
                 location: str = "Zone A") -> ConsolidatedProject:
    refs = [
        SourceReference(
            record_id=f"doc_1:0001#00{n}", doc_id="doc_1", document_title="FY2027-2031 CIP",
            url="https://www.examplecity.gov/files/cip-2027.pdf",
            page_url=f"https://www.examplecity.gov/files/cip-2027.pdf#page={10 + n}", pdf_pages=[10 + n],
            printed_pages=[str(n)], excerpt=f"ST8510000{n}  {name}", original_project_name=name,
            original_project_id=f"ST8510000{n}", document_fiscal_year=2027, estimated_budget=budget, verified=True,
        )
    ]
    return ConsolidatedProject(
        project_uid=f"PRJ-000000000{n}", organization_name=ORG, project_id=f"ST8510000{n}", project_name=name,
        domain=domain, location=location, estimated_budget=budget, budget_type="adopted", current_phase=phase,
        funding_status=funding, source_document="FY2027-2031 CIP",
        source_url="https://www.examplecity.gov/files/cip-2027.pdf", source_pages=[10 + n], references=refs,
        extraction_confidence=0.9, merged_record_ids=[r.record_id for r in refs],
    )


def canned_result(job_id: str, organization: str) -> JobResult:
    projects = [
        make_project(1, "Harbor Road Resurfacing", "transportation", 1_250_000, Phase.FUNDED, FundingStatus.FUNDED),
        make_project(2, "Cedar Water Main Upgrade", "water", 5_000_000, Phase.DESIGN, FundingStatus.PARTIALLY_FUNDED, "Zone B"),
        make_project(3, "Ridge Park Renovation", "parks_recreation", 300_000, Phase.CONSTRUCTION, FundingStatus.FUNDED, "Zone C"),
    ]
    records = [
        ExtractedRecord(
            record_id=f"doc_1:0001#00{i}", organization_name=organization, project_name=p.project_name,
            project_id=p.project_id, estimated_budget=p.estimated_budget, source_document="FY2027-2031 CIP",
            source_url="https://www.examplecity.gov/files/cip-2027.pdf", source_pages=[10 + i], doc_id="doc_1",
            chunk_id="doc_1:0001",
        )
        for i, p in enumerate(projects, start=1)
    ] + [
        ExtractedRecord(
            record_id="doc_1:0002#000", organization_name=organization, project_name="Harbor Rd Resurfacing",
            project_id="ST85100001", source_document="FY2027-2031 CIP",
            source_url="https://www.examplecity.gov/files/cip-2027.pdf", source_pages=[40], doc_id="doc_1",
            chunk_id="doc_1:0002",
        )
    ]
    docs = make_documents()
    usage = [UsageRecord(stage="extraction", task="extraction", provider="anthropic", model="claude-haiku-5-5",
                         input_tokens=4000, output_tokens=900, cost_usd=0.0123)]
    return JobResult(
        job_id=job_id, organization_name=organization, org_profile=OrgProfile(input_name=organization, canonical_name=organization),
        selected_documents=docs[:1], discovered_documents=docs, projects=projects, records=records, usage=usage,
        metrics=JobMetrics(documents_discovered=2, documents_selected=1, pages_processed=40, chunks_produced=6,
                           projects_extracted_raw=4, unique_projects=3, duplicates_merged=1, total_cost_usd=0.0123,
                           extraction_model="claude-haiku-5-5", screening_provider="heuristic"),
    )


class Env:
    def __init__(self, client: TestClient, manager: JobManager, calls: list) -> None:
        self.client, self.manager, self.calls = client, manager, calls

    def submit(self, org: str = ORG, headers=None, **body):
        return self.client.post(EXTRACT, json={"organization_name": org, **body}, headers=headers or KEY)

    def wait(self, job_id: str, timeout: float = 20.0) -> dict:
        deadline = time.time() + timeout
        while time.time() < deadline:
            body = self.client.get(f"/api/v1/jobs/{job_id}", headers=KEY).json()
            if body["status"] in ("completed", "failed"):
                return body
            time.sleep(0.05)
        raise AssertionError("job did not finish in time")

    def completed_job(self, org: str = ORG) -> str:
        job_id = self.submit(org).json()["job_id"]
        assert self.wait(job_id)["status"] == "completed"
        return job_id


@pytest.fixture
def env(tmp_path, monkeypatch, app_config):
    calls: list = []

    async def fake_run(self, organization, options, job_id, ledger, progress):
        calls.append((organization, options))
        progress("extracting", {"chunks_extracted": 1})
        return canned_result(job_id, organization)

    monkeypatch.setattr(Pipeline, "run", fake_run)
    settings = Settings(data_dir=tmp_path / "data", api_keys="test-key, other-key")
    db = Database(f"sqlite:///{tmp_path / 'jobs.db'}")
    manager = JobManager(settings, app_config, db, FileCache(tmp_path / "cache"))
    app = create_app(settings, app_config, manager)
    with TestClient(app) as client:
        yield Env(client, manager, calls)


class TestAuthAndErrors:
    def test_health_needs_no_key(self, env):
        body = env.client.get("/health").json()
        assert body["status"] == "ok" and body["version"]

    @pytest.mark.parametrize("headers", [{}, {"X-API-Key": "wrong"}, {"X-API-Key": ""}])
    def test_missing_or_wrong_key_is_401_with_structured_body(self, env, headers):
        response = env.client.post(EXTRACT, json={"organization_name": ORG}, headers=headers)
        assert response.status_code == 401
        assert response.json()["error"]["code"] == "unauthorized"
        assert env.client.get("/api/v1/jobs", headers=headers).status_code == 401
        assert env.calls == []

    def test_every_configured_key_is_accepted(self, env):
        assert env.client.get("/api/v1/jobs", headers={"X-API-Key": "other-key"}).status_code == 200

    @pytest.mark.parametrize("body", [{"organization_name": "ab"}, {"organization_name": ORG, "max_documents": 9},
                                      {"organization_name": ORG, "screening": "magic"}, {}])
    def test_invalid_requests_are_422(self, env, body):
        response = env.client.post(EXTRACT, json=body, headers=KEY)
        assert response.status_code == 422
        assert response.json()["error"]["code"] == "invalid_request"

    def test_unknown_job_is_404(self, env):
        for suffix in ("", "/projects", "/sources", "/metrics", "/records", "/export.xlsx"):
            response = env.client.get(f"/api/v1/jobs/job_nope{suffix}", headers=KEY)
            assert response.status_code == 404, suffix
            assert response.json()["error"]["code"] == "job_not_found"

    def test_results_of_an_unfinished_job_are_409(self, env):
        env.manager.db.create_job(JobRow(id="job_running", request_key="k", organization_name=ORG, params={},
                                         status="running", stage="extracting", progress={}, warnings=[], errors=[]))
        for suffix in ("/projects", "/sources", "/metrics", "/export.xlsx"):
            response = env.client.get(f"/api/v1/jobs/job_running{suffix}", headers=KEY)
            assert response.status_code == 409, suffix
            error = response.json()["error"]
            assert error["code"] == "job_not_completed" and error["details"]["status"] == "running"


class TestSubmission:
    def test_post_accepts_then_job_completes(self, env):
        response = env.submit()
        assert response.status_code == 202
        body = response.json()
        assert body["created"] is True and body["job_id"].startswith("job_")
        assert body["links"]["projects"] == f"/api/v1/jobs/{body['job_id']}/projects"
        status = env.wait(body["job_id"])
        assert status["status"] == "completed" and status["stage"] == "completed"
        assert status["organization_name"] == ORG and status["attempts"] == 1
        assert status["progress"]["unique_projects"] == 3
        assert status["accumulated_cost_usd"] == pytest.approx(0.0123)
        assert status["finished_at"] and status["errors"] == []
        assert len(env.calls) == 1

    def test_same_request_again_is_idempotent(self, env):
        first = env.submit().json()
        second = env.submit()
        assert second.status_code == 200
        assert second.json()["created"] is False and second.json()["job_id"] == first["job_id"]
        env.wait(first["job_id"])
        third = env.submit()  # still reused after completion (within the reuse window)
        assert third.status_code == 200 and third.json()["job_id"] == first["job_id"]
        assert len(env.calls) == 1

    def test_force_refresh_creates_a_new_job(self, env):
        first = env.submit().json()
        env.wait(first["job_id"])
        again = env.submit(force_refresh=True)
        assert again.status_code == 202 and again.json()["job_id"] != first["job_id"]

    def test_idempotency_key_reuse_with_same_body_returns_the_job(self, env):
        headers = {**KEY, "Idempotency-Key": "abc-123"}
        first = env.submit(headers=headers)
        again = env.submit(headers=headers)
        assert first.status_code == 202 and again.status_code == 200
        assert again.json()["job_id"] == first.json()["job_id"] and again.json()["created"] is False

    def test_idempotency_key_reuse_with_a_different_body_is_409(self, env):
        headers = {**KEY, "Idempotency-Key": "abc-123"}
        env.submit(headers=headers)
        clash = env.submit("Another Organization", headers=headers)
        assert clash.status_code == 409
        assert clash.json()["error"]["code"] == "idempotency_conflict"

    def test_job_list_contains_the_job(self, env):
        job_id = env.submit().json()["job_id"]
        listing = env.client.get("/api/v1/jobs?limit=5", headers=KEY).json()
        assert [j["job_id"] for j in listing] == [job_id]

    def test_failed_job_reports_a_structured_error(self, env, monkeypatch):
        async def boom(self, *args, **kwargs):
            raise RuntimeError("kaboom")

        monkeypatch.setattr(Pipeline, "run", boom)
        status = env.wait(env.submit().json()["job_id"])
        assert status["status"] == "failed"
        assert status["errors"][0]["code"] == "internal_error" and "kaboom" in status["errors"][0]["message"]
        assert env.client.get(f"/api/v1/jobs/{status['job_id']}/projects", headers=KEY).status_code == 409

    def test_missing_provider_key_is_reported_as_such(self, env, monkeypatch):
        async def unavailable(self, *args, **kwargs):
            raise LLMUnavailableError("ANTHROPIC_API_KEY is not set")

        monkeypatch.setattr(Pipeline, "run", unavailable)
        status = env.wait(env.submit().json()["job_id"])
        assert status["status"] == "failed" and status["errors"][0]["code"] == "provider_not_configured"


class TestResults:
    def projects(self, env, job_id, **params):
        response = env.client.get(f"/api/v1/jobs/{job_id}/projects", params=params, headers=KEY)
        assert response.status_code == 200, response.text
        return response.json()

    def test_pagination(self, env):
        job_id = env.completed_job()
        page1 = self.projects(env, job_id, page_size=2)
        assert page1["total"] == 3 and len(page1["items"]) == 2 and page1["next_page"] == 2
        assert page1["organization_name"] == ORG and page1["page"] == 1
        page2 = self.projects(env, job_id, page_size=2, page=2)
        assert len(page2["items"]) == 1 and page2["next_page"] is None
        names = [p["project_name"] for p in page1["items"] + page2["items"]]
        assert sorted(names) == ["Cedar Water Main Upgrade", "Harbor Road Resurfacing", "Ridge Park Renovation"]
        assert env.client.get(f"/api/v1/jobs/{job_id}/projects?page_size=0", headers=KEY).status_code == 422

    def test_filters(self, env):
        job_id = env.completed_job()
        name = lambda **kw: [p["project_name"] for p in self.projects(env, job_id, **kw)["items"]]  # noqa: E731
        assert name(phase="Funded") == ["Harbor Road Resurfacing"]
        assert name(phase="construction / in progress") == ["Ridge Park Renovation"]
        assert name(domain="water") == ["Cedar Water Main Upgrade"]
        assert sorted(name(min_budget=1_000_000)) == ["Cedar Water Main Upgrade", "Harbor Road Resurfacing"]
        assert name(min_budget=10_000_000) == []
        assert name(q="PARK") == ["Ridge Park Renovation"]
        assert name(funding_status="Partially funded") == ["Cedar Water Main Upgrade"]
        assert self.projects(env, job_id, domain="water", min_budget=9e9)["total"] == 0

    def test_project_items_carry_page_level_references(self, env):
        job_id = env.completed_job()
        item = self.projects(env, job_id)["items"][0]
        ref = item["references"][0]
        assert "#page=" in ref["page_url"] and ref["pdf_pages"] and ref["doc_id"] == "doc_1"
        assert item["current_phase"] in {"Funded", "Design", "Construction / In Progress"}

    def test_sources_metrics_and_records(self, env):
        job_id = env.completed_job()
        sources = env.client.get(f"/api/v1/jobs/{job_id}/sources", headers=KEY).json()
        assert [d["doc_id"] for d in sources["selected"]] == ["doc_1"]
        assert len(sources["discovered"]) == 2 and sources["selected"][0]["selection_rationale"]
        metrics = env.client.get(f"/api/v1/jobs/{job_id}/metrics", headers=KEY).json()
        assert metrics["metrics"]["unique_projects"] == 3 and metrics["metrics"]["total_cost_usd"] == pytest.approx(0.0123)
        records = env.client.get(f"/api/v1/jobs/{job_id}/records?page_size=3", headers=KEY).json()
        assert records["total"] == 4 and len(records["items"]) == 3 and records["next_page"] == 2

    def test_excel_export_is_a_valid_workbook_with_page_links(self, env):
        job_id = env.completed_job()
        response = env.client.get(f"/api/v1/jobs/{job_id}/export.xlsx", headers=KEY)
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
        assert "capital-projects-city-of-exampleville.xlsx" in response.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(response.content))
        assert wb.sheetnames[0] == "Projects"
        for sheet in ("Sources", "Project References", "Extraction & Cost Metrics"):
            assert sheet in wb.sheetnames
        projects = wb["Projects"]
        texts = {str(c.value) for row in projects.iter_rows() for c in row if c.value is not None}
        assert any("budget" in t.lower() for t in texts)
        assert {"Harbor Road Resurfacing", "Cedar Water Main Upgrade", "Ridge Park Renovation"} <= texts
        # xlsxwriter stores 'file.pdf#page=12' as target 'file.pdf' + location 'page=12'
        links = [(c.hyperlink.target or "") + (f"#{c.hyperlink.location}" if c.hyperlink.location else "")
                 for ws in wb for row in ws.iter_rows() for c in row if c.hyperlink]
        assert any("#page=" in link for link in links)

    def test_openapi_lists_the_endpoints(self, env):
        paths = env.client.get("/openapi.json").json()["paths"]
        for path in ("/health", EXTRACT, "/api/v1/jobs", "/api/v1/jobs/{job_id}", "/api/v1/jobs/{job_id}/projects",
                     "/api/v1/jobs/{job_id}/records", "/api/v1/jobs/{job_id}/sources", "/api/v1/jobs/{job_id}/metrics",
                     "/api/v1/jobs/{job_id}/export.xlsx"):
            assert path in paths, path
        assert env.client.get("/docs").status_code == 200


class TestRegression:
    def test_valid_key_gets_a_2xx_on_authenticated_endpoints(self, env):
        # regression: FastAPI used to treat the dependency parameters `_` and `manager` as required query
        # parameters (function-local aliases + postponed annotations), so every authenticated call returned 422
        assert env.client.get("/api/v1/jobs", headers=KEY).status_code == 200
        assert env.submit().status_code == 202
        assert env.client.get("/api/v1/jobs").status_code == 401
