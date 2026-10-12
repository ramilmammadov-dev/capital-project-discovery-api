"""Offline end-to-end run of the whole pipeline.

search (DuckDuckGo, mocked) -> official landing page crawl -> two PDFs (an itemized CIP and an aggregate
budget summary) -> probe -> rank/select -> parse -> chunk -> heuristic screening -> extraction (fake LLM that
answers from the chunk text) -> validation -> dedup -> Excel. A second run proves that repeat runs make no new
model calls.
"""

from __future__ import annotations

from pathlib import Path

import httpx
import openpyxl
import pytest
import respx

from capex.config import AppConfig
from capex.costs import CostLedger
from capex.export import save_outputs
from capex.llm import LLMRegistry
from capex.pipeline import Pipeline, PipelineOptions
from capex.storage import Database, FileCache
from tests.helpers import FakeLLM, build_cip_pdf, build_summary_pdf, llm_ok, projects_from_user

SITE = "https://www.examplecity.gov"
CIP_URL = f"{SITE}/files/cip-fy2027-2031.pdf"
SUMMARY_URL = f"{SITE}/files/budget-summary-fy2027.pdf"

DDG_PAGE = f"""<html><body>
<div class="result"><a class="result__a" href="{SITE}/cip">Capital Improvement Program | City of Exampleville</a>
<a class="result__snippet">The adopted five-year capital improvement program lists every project.</a></div>
</body></html>"""

LANDING = f"""<html><head><title>Capital Improvement Program</title></head><body><main>
<h1>Capital Improvement Program</h1>
<p>The City adopts a five-year capital improvement program each June.</p>
<ul>
  <li><a href="{CIP_URL}">FY 2027-2031 Capital Improvement Program (Adopted)</a></li>
  <li><a href="{SUMMARY_URL}">FY 2027 Annual Budget Summary</a></li>
</ul></main></body></html>"""


def _handler(call: dict) -> object:
    name = call["schema_name"]
    if name == "org_search_plan":
        return llm_ok(
            {
                "canonical_name": "City of Exampleville",
                "org_type": "city",
                "state": "AZ",
                "official_domains": ["examplecity.gov"],
                "aliases": ["Exampleville"],
                "document_terms": ["Capital Improvement Program"],
                "queries": ["City of Exampleville capital improvement program"],
            }
        )
    if name == "document_profiles":
        return llm_ok({"documents": []})
    if name == "capital_projects":
        return llm_ok(projects_from_user(call["user"]))
    raise AssertionError(f"unexpected LLM call {name}")


@pytest.fixture
def site(tmp_path: Path):
    cip_meta = build_cip_pdf(tmp_path / "cip.pdf", n_rows=70)
    build_summary_pdf(tmp_path / "summary.pdf")
    with respx.mock(assert_all_called=False) as router:
        router.post("https://html.duckduckgo.com/html/").mock(return_value=httpx.Response(200, text=DDG_PAGE))
        router.get(f"{SITE}/robots.txt").mock(return_value=httpx.Response(404))
        router.get(f"{SITE}/cip").mock(
            return_value=httpx.Response(200, text=LANDING, headers={"content-type": "text/html"})
        )
        router.get(CIP_URL).mock(
            return_value=httpx.Response(200, content=(tmp_path / "cip.pdf").read_bytes(), headers={"content-type": "application/pdf"})
        )
        router.get(SUMMARY_URL).mock(
            return_value=httpx.Response(200, content=(tmp_path / "summary.pdf").read_bytes(), headers={"content-type": "application/pdf"})
        )
        yield router, cip_meta


@pytest.fixture
def env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    for var in ("ANTHROPIC_API_KEY", "TYPESAFE_API_KEY", "OPENAI_API_KEY", "BRAVE_API_KEY", "SERPER_API_KEY"):
        monkeypatch.delenv(var, raising=False)
    config = AppConfig()
    config.search.providers = ["duckduckgo"]
    config.parsing.workers = 1
    config.parsing.ocr_enabled = False
    db = Database(f"sqlite:///{tmp_path / 'capex.db'}")
    cache = FileCache(tmp_path / "cache")
    fake = FakeLLM(_handler)
    registry = LLMRegistry(config)
    registry.register("anthropic", fake)
    return config, db, cache, registry, fake


async def _run(config, db, cache, registry, job_id: str):
    ledger = CostLedger(config)
    pipeline = Pipeline(config, cache, db, registry)
    result = await pipeline.run("City of Exampleville", PipelineOptions(screening="heuristic"), job_id, ledger)
    return result, ledger


async def test_end_to_end_pipeline(site, env, tmp_path: Path) -> None:
    _router, meta = site
    config, db, cache, registry, fake = env

    result, ledger = await _run(config, db, cache, registry, "job_e2e_1")

    # discovery & selection: the itemized CIP is the top document, the aggregate summary is recognised as such
    assert result.selected_documents, result.metrics.warnings
    top = result.selected_documents[0]
    assert top.url == CIP_URL
    assert top.selection_rank == 1 and top.selection_rationale
    summary = next(d for d in result.discovered_documents if d.url == SUMMARY_URL)
    assert summary.contains_itemized_projects is False

    # every project row of the CIP becomes exactly one unique project, cited to the page it is printed on
    ids = set(meta["ids"])
    found = {p.project_id for p in result.projects}
    assert ids <= found, sorted(ids - found)[:5]
    assert len([p for p in result.projects if p.project_id in ids]) == len(ids)
    for record in result.records:
        if record.project_id in meta["page_of"]:
            assert meta["page_of"][record.project_id] in record.source_pages

    m = result.metrics
    assert m.projects_extracted_raw >= m.unique_projects >= len(ids)
    assert m.chunks_produced >= m.chunks_accepted > 0
    assert m.screening_provider.startswith("heuristic")
    assert "extraction" in m.cost_by_stage and m.total_cost_usd > 0
    assert all(r.cost_usd >= 0 for r in ledger.records)

    # deliverable workbook
    paths = save_outputs(result, tmp_path / "out")
    wb = openpyxl.load_workbook(paths["xlsx"])
    assert wb.sheetnames[:4] == ["Projects", "Sources", "Project References", "Extraction & Cost Metrics"]
    assert wb["Projects"].max_row - 1 >= len(ids)

    # repeat run: everything is served from caches -> no new model calls
    extraction_calls = len(fake.calls_for("capital_projects"))
    second, ledger2 = await _run(config, db, cache, registry, "job_e2e_2")
    assert len(fake.calls_for("capital_projects")) == extraction_calls
    assert ledger2.total_cost == 0
    # stable project IDs across runs (registry), no duplicate records
    assert {p.project_uid for p in second.projects} == {p.project_uid for p in result.projects}
